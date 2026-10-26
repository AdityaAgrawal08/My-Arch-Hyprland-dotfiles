#!/usr/bin/env python3
"""Create small document fixtures inside the repository state directory."""

from __future__ import annotations

import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
VENV_PY = (
    Path(
        os.environ.get(
            "DOC_TOOLS_VENV",
            Path.home() / ".local" / "share" / "doctools" / "venv",
        )
    )
    / "bin"
    / "python"
)

# Make the script work when invoked directly, without forcing the caller to
# remember the venv path.
if (
    VENV_PY.exists()
    and Path(sys.prefix).resolve() != VENV_PY.parent.parent.resolve()
    and os.environ.get("DOC_TOOLS_FIXTURES_REEXEC") != "1"
):
    os.environ["DOC_TOOLS_FIXTURES_REEXEC"] = "1"
    os.execv(str(VENV_PY), [str(VENV_PY), str(Path(__file__).resolve()), *sys.argv[1:]])

DEFAULT = ROOT / "state" / "doc-tools-fixtures"


def _ensure_inside_repo(path: Path) -> Path:
    path = path.resolve()
    try:
        path.relative_to(ROOT.resolve())
    except ValueError as exc:
        raise SystemExit(f"refusing to write outside {ROOT}: {path}") from exc
    path.mkdir(parents=True, exist_ok=True)
    return path


def make_pdf(path: Path, pages: int, label: str) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for page in range(1, pages + 1):
        c.drawString(72, 720, f"{label} page {page}")
        c.showPage()
    c.save()


def make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Doc-tools sample document", 0)
    document.add_paragraph("This fixture is generated inside the repository.")
    document.add_paragraph("It is used by the docx2pdf and pdf2docx tests.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    document.save(path)


def make_image(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1200, 400), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("/usr/share/fonts/TTF/DejaVuSans.ttf")
    font = ImageFont.truetype(str(font_path), 64) if font_path.exists() else None
    draw.text((60, 100), "Hello OCR 12345", font=font, fill="black")
    image.save(path)


def main() -> int:
    target = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else DEFAULT
    target = _ensure_inside_repo(target)

    make_pdf(target / "alpha.pdf", 2, "alpha")
    make_pdf(target / "beta.pdf", 3, "beta")
    make_pdf(target / "gamma.pdf", 1, "gamma")

    make_pdf(
        target
        / "a-very-long-filename-that-should-not-overflow-the-table-of-contents-line-in-a4-page.pdf",
        1,
        "long",
    )
    make_pdf(target / "résumé.pdf", 1, "resume")
    make_pdf(target / "日本語のファイル.pdf", 1, "cjk")

    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(str(target / "alpha.pdf"))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret", "secret", algorithm="AES-256")
    with open(target / "encrypted.pdf", "wb") as handle:
        writer.write(handle)

    (target / "broken.pdf").write_text("this is not a PDF\n", encoding="utf-8")

    make_docx(target / "sample.docx")
    make_image(target / "scan.png")

    (target / "sample.md").write_text(
        "# Doc-tools Markdown\n\n"
        "This is **bold**, this is *italic*.\n\n"
        "- one\n"
        "- two\n\n"
        "```python\nprint('hello')\n```\n\n"
        "| Key | Value |\n| --- | --- |\n| alpha | 1 |\n",
        encoding="utf-8",
    )

    print(f"fixtures: {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
